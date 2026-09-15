from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re

from docx import Document
from pypdf import PdfReader


class ResumeParseError(ValueError):
    """Raised when an uploaded resume cannot be safely parsed."""


@dataclass(frozen=True)
class ParsedResume:
    text: str
    resume: dict[str, object]


def parse_resume_file(path: Path, suffix: str) -> ParsedResume:
    normalized_suffix = suffix.lower()
    try:
        if normalized_suffix == ".pdf":
            text = _extract_pdf_text(path)
        elif normalized_suffix == ".docx":
            text = _extract_docx_text(path)
        else:
            raise ResumeParseError("仅支持 PDF 和 DOCX 格式的简历文件。")
    except ResumeParseError:
        raise
    except Exception as error:
        raise ResumeParseError("简历文件无法解析，请确认文件未损坏。") from error

    text = "\n".join(line.strip() for line in text.splitlines() if line.strip()).strip()
    if not text:
        raise ResumeParseError("简历文件为空或未包含可提取的文字。")
    return ParsedResume(text=text, resume=_normalize_resume(text))


def _extract_pdf_text(path: Path) -> str:
    reader = PdfReader(str(path))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    return text


def _extract_docx_text(path: Path) -> str:
    document = Document(str(path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(paragraphs)


def _normalize_resume(text: str) -> dict[str, object]:
    resume = {
        "version": 1,
        "basic": {"name": "", "phone": "", "email": "", "city": ""},
        "job": {"target_role": "", "employment_type": "", "expected_salary": ""},
        "education": [],
        "employment": [],
        "projects": [],
        "skills": {"skills": [], "certificates": []},
        "self_evaluation": "",
        "section_visibility": {
            "basic": True,
            "job": True,
            "education": True,
            "employment": True,
            "projects": True,
            "skills": True,
            "self_evaluation": True,
        },
    }
    aliases = {
        "姓名": ("basic", "name"),
        "name": ("basic", "name"),
        "电话": ("basic", "phone"),
        "手机": ("basic", "phone"),
        "phone": ("basic", "phone"),
        "邮箱": ("basic", "email"),
        "电子邮箱": ("basic", "email"),
        "email": ("basic", "email"),
        "城市": ("basic", "city"),
        "所在地": ("basic", "city"),
        "city": ("basic", "city"),
        "目标职位": ("job", "target_role"),
        "求职意向": ("job", "target_role"),
        "目标岗位": ("job", "target_role"),
        "target role": ("job", "target_role"),
    }
    unmatched: list[str] = []
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        match = re.match(r"^([^:：]{1,30})\s*[:：]\s*(.+)$", line)
        key = match.group(1).strip().lower() if match else ""
        value = match.group(2).strip() if match else ""
        target = aliases.get(key)
        if target and value:
            section, field = target
            section_data = resume[section]
            assert isinstance(section_data, dict)
            section_data[field] = value
        else:
            unmatched.append(line)
    resume["self_evaluation"] = "\n".join(unmatched)
    return resume
