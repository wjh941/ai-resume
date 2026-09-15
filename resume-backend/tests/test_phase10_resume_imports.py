from __future__ import annotations

from docx import Document

from conftest import make_draft_payload


def _valid_pdf() -> bytes:
    objects = [
        b"1 0 obj\n<</Type/Catalog/Pages 2 0 R>>\nendobj\n",
        b"2 0 obj\n<</Type/Pages/Kids[3 0 R]/Count 1>>\nendobj\n",
        b"3 0 obj\n<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 300]/Contents 4 0 R/Resources<</Font<</F1<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>>>>>\nendobj\n",
        b"4 0 obj\n<</Length 72>>\nstream\nBT /F1 12 Tf 20 250 Td (Name: Alice) Tj 0 -20 Td (Email: alice@example.com) Tj ET\nendstream\nendobj\n",
    ]
    document = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for obj in objects:
        offsets.append(len(document))
        document.extend(obj)
    xref_offset = len(document)
    document.extend(b"xref\n0 5\n0000000000 65535 f \n")
    document.extend(b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:]))
    document.extend(b"trailer\n<</Size 5/Root 1 0 R>>\nstartxref\n")
    document.extend(f"{xref_offset}\n%%EOF\n".encode())
    return bytes(document)


def _valid_docx() -> bytes:
    from io import BytesIO

    buffer = BytesIO()
    document = Document()
    document.add_paragraph("姓名：张三")
    document.add_paragraph("邮箱：zhangsan@example.com")
    document.save(buffer)
    return buffer.getvalue()


def _saved_draft(api_client) -> str:
    response = api_client.post("/api/draft/save", json=make_draft_payload())
    assert response.status_code == 200
    return response.json()["data"]["id"]


def test_pdf_upload_returns_parsed_resume_preview(api_client) -> None:
    draft_id = _saved_draft(api_client)

    response = api_client.post(
        f"/api/draft/{draft_id}/imports",
        files={"file": ("resume.pdf", _valid_pdf(), "application/pdf")},
    )

    assert response.status_code == 200
    data = response.json()["data"]
    assert data["status"] == "parsed"
    assert data["original_filename"] == "resume.pdf"
    assert data["parsed_resume"]["basic"]["name"] == "Alice"
    assert data["parsed_resume"]["basic"]["email"] == "alice@example.com"
    assert "path" not in data


def test_upload_rejects_unsafe_type(api_client) -> None:
    draft_id = _saved_draft(api_client)

    response = api_client.post(
        f"/api/draft/{draft_id}/imports",
        files={"file": ("resume.exe", b"MZ", "application/octet-stream")},
    )

    assert response.status_code == 422
    assert response.json()["code"] == "validation_error"
    assert response.json()["message"] == "仅支持 PDF 和 DOCX 格式的简历文件。"


def test_upload_rejects_doc_extension(api_client) -> None:
    draft_id = _saved_draft(api_client)

    response = api_client.post(
        f"/api/draft/{draft_id}/imports",
        files={"file": ("resume.doc", b"legacy document", "application/msword")},
    )

    assert response.status_code == 422
    assert response.json()["message"] == "仅支持 PDF 和 DOCX 格式的简历文件。"


def test_upload_rejects_pdf_with_invalid_signature_and_removes_file(api_client) -> None:
    draft_id = _saved_draft(api_client)
    import_dir = api_client.app.state.settings.temp_file_path / "resume-imports"

    response = api_client.post(
        f"/api/draft/{draft_id}/imports",
        files={"file": ("resume.pdf", b"not a pdf", "application/pdf")},
    )

    assert response.status_code == 422
    assert response.json()["message"] == "简历文件格式无效，请上传真实的 PDF 或 DOCX 文件。"
    assert not import_dir.exists() or not list(import_dir.iterdir())


def test_upload_rejects_docx_with_invalid_signature_and_removes_file(api_client) -> None:
    draft_id = _saved_draft(api_client)
    import_dir = api_client.app.state.settings.temp_file_path / "resume-imports"

    response = api_client.post(
        f"/api/draft/{draft_id}/imports",
        files={
            "file": (
                "resume.docx",
                b"not a docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 422
    assert response.json()["message"] == "简历文件格式无效，请上传真实的 PDF 或 DOCX 文件。"
    assert not import_dir.exists() or not list(import_dir.iterdir())


def test_docx_upload_returns_parsed_resume_preview(api_client) -> None:
    draft_id = _saved_draft(api_client)

    response = api_client.post(
        f"/api/draft/{draft_id}/imports",
        files={
            "file": (
                "resume.docx",
                _valid_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    data = response.json()["data"]
    assert data["status"] == "parsed"
    assert data["parsed_resume"]["basic"]["name"] == "张三"


def test_upload_rejects_file_larger_than_configured_limit(api_client) -> None:
    draft_id = _saved_draft(api_client)
    oversized = b"x" * (api_client.app.state.settings.resume_import_max_file_bytes + 1)

    response = api_client.post(
        f"/api/draft/{draft_id}/imports",
        files={"file": ("resume.docx", oversized, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )

    assert response.status_code == 422
    assert response.json()["code"] == "validation_error"
    assert response.json()["message"] == "简历文件超过当前允许的大小限制。"
