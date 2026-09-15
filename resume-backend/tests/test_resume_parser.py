from __future__ import annotations

from docx import Document
import pytest

from app.services.resume_parser import ResumeParseError, parse_resume_file


def test_parse_docx_extracts_labeled_fields_and_preserves_unmatched_text(tmp_path):
    path = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("姓名：张三")
    document.add_paragraph("电话: 13800138000")
    document.add_paragraph("邮箱：zhangsan@example.com")
    document.add_paragraph("城市：上海")
    document.add_paragraph("目标职位：后端工程师")
    document.add_paragraph("熟悉 Python、FastAPI")
    document.save(path)

    result = parse_resume_file(path, ".docx")

    assert "熟悉 Python、FastAPI" in result.text
    assert result.resume["basic"] == {
        "name": "张三",
        "phone": "13800138000",
        "email": "zhangsan@example.com",
        "city": "上海",
    }
    assert result.resume["job"]["target_role"] == "后端工程师"
    assert result.resume["self_evaluation"] == "熟悉 Python、FastAPI"


def test_parse_pdf_extracts_text_and_labeled_fields(tmp_path):
    path = tmp_path / "resume.pdf"
    objects = [
        b"1 0 obj\n<</Type/Catalog/Pages 2 0 R>>\nendobj\n",
        b"2 0 obj\n<</Type/Pages/Kids[3 0 R]/Count 1>>\nendobj\n",
        b"3 0 obj\n<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 300]/Contents 4 0 R/Resources<</Font<</F1<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>>>>>\nendobj\n",
        b"4 0 obj\n<</Length 79>>\nstream\nBT /F1 12 Tf 20 250 Td (Name: Alice) Tj 0 -20 Td (Email: alice@example.com) Tj ET\nendstream\nendobj\n",
    ]
    document = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for obj in objects:
        offsets.append(len(document))
        document.extend(obj)
    xref_offset = len(document)
    document.extend(b"xref\n0 5\n0000000000 65535 f \n")
    document.extend(b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:]))
    document.extend(b"trailer\n<</Size 5/Root 1 0 R>>\nstartxref\n")
    document.extend(f"{xref_offset}\n%%EOF\n".encode())
    path.write_bytes(document)

    result = parse_resume_file(path, ".pdf")

    assert "Alice" in result.text
    assert result.resume["basic"]["name"] == "Alice"
    assert result.resume["basic"]["email"] == "alice@example.com"


@pytest.mark.parametrize("suffix", [".pdf", ".docx"])
def test_empty_or_malformed_document_raises_parser_error(tmp_path, suffix):
    path = tmp_path / f"broken{suffix}"
    path.write_bytes(b"not a document")

    with pytest.raises(ResumeParseError):
        parse_resume_file(path, suffix)


def test_empty_docx_raises_parser_error(tmp_path):
    path = tmp_path / "empty.docx"
    Document().save(path)

    with pytest.raises(ResumeParseError):
        parse_resume_file(path, ".docx")


def test_unsupported_suffix_raises_parser_error(tmp_path):
    path = tmp_path / "resume.txt"
    path.write_text("姓名：张三", encoding="utf-8")

    with pytest.raises(ResumeParseError):
        parse_resume_file(path, ".txt")
